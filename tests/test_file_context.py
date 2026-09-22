# Kestrel - AI desktop automation for Windows.
# Copyright (C) 2026 the Kestrel authors
# Dual-licensed: GNU GPL v3 (see LICENSE), or a commercial licence.
# This program comes with ABSOLUTELY NO WARRANTY.

"""
tests/test_file_context.py
============================
Tests for ai/file_context.py - extracting text from files dropped into
the chat input.
"""

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import ai.file_context as file_context


def test_extract_text_from_txt_file(tmp_path):
    f = tmp_path / "notes.txt"
    f.write_text("hello world", encoding="utf-8")

    text, err = file_context.extract_text_from_file(str(f))

    assert err is None
    assert text == "hello world"


def test_extract_text_from_json_file(tmp_path):
    f = tmp_path / "data.json"
    f.write_text('{"key": "value"}', encoding="utf-8")

    text, err = file_context.extract_text_from_file(str(f))

    assert err is None
    assert text == '{"key": "value"}'


def test_extract_text_from_md_csv_log_files(tmp_path):
    for ext in (".md", ".csv", ".log"):
        f = tmp_path / f"file{ext}"
        f.write_text("some content", encoding="utf-8")
        text, err = file_context.extract_text_from_file(str(f))
        assert err is None
        assert text == "some content"


def test_extract_text_rejects_unsupported_extension(tmp_path):
    f = tmp_path / "video.mp4"
    f.write_bytes(b"binary data")

    text, err = file_context.extract_text_from_file(str(f))

    assert text is None
    assert "Unsupported file type" in err
    assert ".mp4" in err


def test_extract_text_reports_missing_file():
    text, err = file_context.extract_text_from_file("/nonexistent/path/file.txt")

    assert text is None
    assert "not found" in err.lower()


def test_extract_text_truncates_long_files(tmp_path):
    f = tmp_path / "big.txt"
    f.write_text("x" * (file_context.MAX_FILE_CONTEXT_CHARS + 500), encoding="utf-8")

    text, err = file_context.extract_text_from_file(str(f))

    assert err is None
    assert len(text) <= file_context.MAX_FILE_CONTEXT_CHARS + len("\n... (truncated)")
    assert text.endswith("... (truncated)")


def test_extract_text_handles_non_utf8_gracefully(tmp_path):
    f = tmp_path / "weird.txt"
    f.write_bytes(b"\xff\xfe some garbage bytes")

    text, err = file_context.extract_text_from_file(str(f))

    assert err is None  # errors="replace" - never raises on bad encoding
    assert text is not None


def test_extract_pdf_text_reports_missing_pypdf(tmp_path, monkeypatch):
    f = tmp_path / "doc.pdf"
    f.write_bytes(b"%PDF-1.4 fake content")

    import builtins
    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "pypdf":
            raise ImportError("no pypdf")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)

    text, err = file_context.extract_text_from_file(str(f))

    assert text is None
    assert "pypdf" in err


def test_format_attachment_block():
    block = file_context.format_attachment_block("report.txt", "Q3 revenue is up.")
    assert block == "[Attached file: report.txt]\nQ3 revenue is up."


# ══════════════════════════════════════════════════════════════════
# .docx / .xlsx - added alongside tool_read_file / tool_merge_files_to_pdf
# ══════════════════════════════════════════════════════════════════

def test_extract_docx_text_paragraphs_and_tables(tmp_path):
    import docx
    document = docx.Document()
    document.add_paragraph("Hello from a paragraph.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "Score"
    f = tmp_path / "report.docx"
    document.save(str(f))

    text, err = file_context.extract_text_from_file(str(f))

    assert err is None
    assert "Hello from a paragraph." in text
    assert "Name" in text and "Score" in text


def test_extract_docx_text_reports_empty_document(tmp_path):
    import docx
    document = docx.Document()  # no paragraphs, no tables
    f = tmp_path / "empty.docx"
    document.save(str(f))

    text, err = file_context.extract_text_from_file(str(f))

    assert text is None
    assert "No extractable text" in err


def test_extract_docx_text_reports_missing_python_docx(tmp_path, monkeypatch):
    f = tmp_path / "doc.docx"
    f.write_bytes(b"fake docx bytes")

    import builtins
    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "docx":
            raise ImportError("no docx")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)

    text, err = file_context.extract_text_from_file(str(f))

    assert text is None
    assert "python-docx" in err


def test_extract_xlsx_text_reads_sheet_data(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Score"])
    ws.append(["Alice", 90])
    f = tmp_path / "grades.xlsx"
    wb.save(str(f))

    text, err = file_context.extract_text_from_file(str(f))

    assert err is None
    assert "Sheet1" in text
    assert "Name" in text and "Alice" in text and "90" in text


def test_extract_xlsx_text_reports_empty_workbook(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()  # default empty sheet, no data
    f = tmp_path / "empty.xlsx"
    wb.save(str(f))

    text, err = file_context.extract_text_from_file(str(f))

    assert text is None
    assert "No data found" in err


def test_extract_xlsx_text_reports_missing_openpyxl(tmp_path, monkeypatch):
    f = tmp_path / "sheet.xlsx"
    f.write_bytes(b"fake xlsx bytes")

    import builtins
    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "openpyxl":
            raise ImportError("no openpyxl")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)

    text, err = file_context.extract_text_from_file(str(f))

    assert text is None
    assert "openpyxl" in err


def test_supported_extensions_include_docx_and_xlsx():
    assert ".docx" in file_context.SUPPORTED_EXTENSIONS
    assert ".xlsx" in file_context.SUPPORTED_EXTENSIONS
